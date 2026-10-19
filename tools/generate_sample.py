"""예시 과제 데이터 생성기 (가상의 반도체 제조 조직: 6개 팀 × 2개 그룹, 36개 과제).

실행: python tools/generate_sample.py   (openpyxl 필요: pip install openpyxl)
결과: sample/tasks.json, sample/history.json, sample/AX_과제_예시.xlsx,
      index.html 안의 예시 데이터(<script id="sample-data">, <script id="sample-history">)
"""
import json
import random
import re as regex
import datetime as dt
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
random.seed(23)
ASOF = dt.date(2026, 10, 19)
U = random.uniform
D = lambda d: d.isoformat() if d else None
r2 = lambda x: None if x is None else round(x, 2)

# 팀 → 그룹 → (대분류, 중분류, 구분, Agent, 지표, KPI, 단위, 방향, AS-IS, TO-BE, 적용범위, 라인/공정)
ORG = {
    "제조기술팀": {
        "공정기술그룹": [
            ("공정 최적화", "AI 공정 레시피 추천", "AX", "O", "Q", "공정 불량률", "%", "감소", 2.4, 1.1, "복수 라인", "FAB1 포토·식각"),
            ("공정 최적화", "공정 조건 이상 탐지", "AX", "X", "Q", "이상 검출 누락", "건", "감소", 22, 5, "공장 전체", "FAB1~2 전 공정"),
            ("공정 최적화", "공정 변경점 문서 자동 작성 Agent", "AX", "O", "P", "공수 절감", "FTE", None, None, None, "팀 업무", "공정기술"),
        ],
        "공정자동화그룹": [
            ("반송 자동화", "AMHS 반송 경로 최적화", "FA", "X", "D", "반송 대기시간", "분", "감소", 14, 6, "공장 전체", "FAB1 AMHS"),
            ("반송 자동화", "로트 자동 디스패칭", "FA", "X", "P", "공수 절감", "FTE", None, None, None, "복수 라인", "FAB2 확산·증착"),
            ("설비 자동화", "설비 셋업 자동화", "FA", "X", "D", "설비 셋업시간", "분", "감소", 50, 20, "단일 라인", "FAB2 증착 3호기"),
        ],
    },
    "품질보증팀": {
        "계측그룹": [
            ("계측 자동화", "계측 데이터 자동 판정", "AX", "X", "P", "공수 절감", "FTE", None, None, None, "팀 업무", "CD·두께 계측"),
            ("계측 자동화", "결함 이미지 자동 분류(ADC)", "AX", "X", "Q", "결함 분류 정확도", "%", "증가", 82, 95, "공장 전체", "결함 리뷰"),
            ("계측 자동화", "계측 샘플링 최적화", "AX", "X", "C", "계측 비용", "백만원", "감소", 380, 300, "복수 라인", "FAB1 계측"),
        ],
        "품질분석그룹": [
            ("품질 분석", "불량 원인분석 Agent", "AX", "O", "Q", "원인분석 소요시간", "시간", "감소", 16, 4, "팀 업무", "품질분석"),
            ("품질 분석", "고객 품질 이슈 자동 분류", "AX", "O", "P", "공수 절감", "FTE", None, None, None, "팀 업무", "고객 품질"),
            ("품질 분석", "SPC 이상 자동 알림", "AX", "X", "Q", "공정능력지수", "Cpk", "증가", 1.1, 1.5, "공장 전체", "전 공정"),
        ],
    },
    "FAB운영1팀": {
        "포토그룹": [
            ("포토 공정", "노광 오버레이 자동 보정", "AX", "X", "Q", "오버레이 불량률", "%", "감소", 1.8, 0.7, "단일 라인", "FAB1 포토"),
            ("포토 공정", "레티클 반출입 자동화", "FA", "X", "P", "공수 절감", "FTE", None, None, None, "단일 라인", "FAB1 레티클 스토커"),
            ("생산 실적", "포토 작업일보 자동 작성 Agent", "AX", "O", "M", "반복업무 비중", "%", "감소", 35, 15, "팀 업무", "포토 생산관리"),
        ],
        "식각그룹": [
            ("식각 공정", "식각 종말점 예측", "AX", "X", "Q", "식각 재작업률", "%", "감소", 3.2, 1.2, "단일 라인", "FAB1 식각"),
            ("식각 공정", "챔버 세정 주기 최적화", "AX", "X", "C", "챔버 소모품비", "백만원", "감소", 520, 430, "복수 라인", "FAB1 식각 챔버"),
            ("식각 공정", "식각 설비 알람 대응 Agent", "AX", "O", "P", "공수 절감", "FTE", None, None, None, "팀 업무", "식각 설비"),
        ],
    },
    "FAB운영2팀": {
        "증착그룹": [
            ("증착 공정", "박막 두께 가상계측", "AX", "X", "Q", "두께 산포", "Å", "감소", 18, 9, "복수 라인", "FAB2 CVD"),
            ("증착 공정", "증착 레시피 자동 세팅", "FA", "X", "D", "레시피 셋업시간", "분", "감소", 45, 15, "단일 라인", "FAB2 CVD 2호기"),
            ("증착 공정", "증착 설비 PM 이력 요약 Agent", "AX", "O", "P", "공수 절감", "FTE", None, None, None, "팀 업무", "증착 설비"),
        ],
        "세정그룹": [
            ("세정 공정", "약액 사용량 최적화", "AX", "X", "C", "약액비", "백만원", "감소", 640, 540, "복수 라인", "FAB2 세정"),
            ("세정 공정", "세정 설비 자동 투입", "FA", "X", "P", "공수 절감", "FTE", None, None, None, "단일 라인", "FAB2 세정 1라인"),
            ("세정 공정", "약액 공급실 안전 모니터링", "FA", "X", "S", "약액 노출 위험작업", "건", "감소", 12, 3, "단일 라인", "FAB2 약액 공급실"),
        ],
    },
    "설비기술팀": {
        "예방보전그룹": [
            ("예지보전", "진동 기반 고장 예측", "AX", "X", "D", "설비 비가동", "시간", "감소", 320, 150, "공장 전체", "주요 설비 42대"),
            ("예지보전", "보전 이력 검색 Agent", "AX", "O", "P", "공수 절감", "FTE", None, None, None, "팀 업무", "보전 이력"),
            ("예지보전", "PM 일정 자동 편성", "AX", "X", "P", "공수 절감", "FTE", None, None, None, "팀 업무", "PM 계획"),
        ],
        "설비개선그룹": [
            ("설비 관리", "스마트 점검 앱", "FA", "X", "S", "점검 누락", "건", "감소", 30, 5, "공장 전체", "일상점검"),
            ("설비 관리", "예비품 재고 최적화", "AX", "X", "C", "예비품 재고", "백만원", "감소", 900, 700, "팀 업무", "예비품 창고"),
            ("설비 관리", "유틸리티 원격 모니터링", "FA", "X", "D", "이상 대응시간", "분", "감소", 40, 10, "공장 전체", "유틸리티동"),
        ],
    },
    "경영지원팀": {
        "구매그룹": [
            ("구매 업무", "발주서 자동 생성 RPA", "AX", "X", "P", "공수 절감", "FTE", None, None, None, "팀 업무", "구매 발주"),
            ("구매 업무", "협력사 납기 회신 Agent", "AX", "O", "D", "납기 준수율", "%", "증가", 88, 97, "팀 업무", "협력사 관리"),
            ("구매 업무", "단가 비교 분석 자동화", "AX", "X", "C", "기타", "시간", "감소", 10, 3, "팀 업무", "구매 단가"),
        ],
        "경영관리그룹": [
            ("인사 업무", "교육이력 자동 집계", "AX", "X", "P", "공수 절감", "FTE", None, None, None, "팀 업무", "교육 관리"),
            ("인사 업무", "사내 규정 Q&A Agent", "AX", "O", "M", "문의 응답시간", "시간", "감소", 24, 2, "팀 업무", "인사 문의"),
            ("원가 관리", "월 결산 자동화", "AX", "O", "P", "공수 절감", "FTE", None, None, None, "팀 업무", "월 결산"),
        ],
    },
}
# 팀별 실적 성향: (확보 배율, 일정 지연 확률). 조직별 현황 화면에서 양호·주의·미흡이 섞여 보이도록 다르게 둠
TEAM_PERF = {"제조기술팀": (1.0, .2), "품질보증팀": (1.05, .1), "FAB운영1팀": (.6, .45), "FAB운영2팀": (.88, .25), "설비기술팀": (1.0, .2), "경영지원팀": (.85, .2)}
DEV = {"AX": ["DX추진팀", "AI솔루션그룹"], "FA": ["스마트팩토리팀", "제조기술팀"]}
NAMES = ["kim.mj", "lee.sh", "park.jy", "choi.dw", "jung.hn", "kang.sb", "yoon.ej", "lim.th", "han.ys", "oh.jk", "seo.hr", "shin.wj"]
STAGE = {1: "후보 발굴·타당성", 2: "개발·구현", 3: "현장 실증·검증", 4: "정식 업무 적용", 5: "표준화·확산"}
ids = lambda: ";".join(random.sample(NAMES, random.choice([1, 1, 2, 3])))

rows, ramps = [], {}
for team, groups in ORG.items():
    mult, p_delay = TEAM_PERF[team]
    for group, specs in groups.items():
        for (cat, task, kind, agent, mg, kpi, unit, dirn, asis, tobe, scope, line) in specs:
            start = dt.date(2025, 5, 12) + dt.timedelta(days=random.randint(0, 430))
            dur = random.randint(300, 560)
            end = start + dt.timedelta(days=dur)
            sd = {k: start + dt.timedelta(days=int(dur * f)) for k, f in zip(range(1, 6), (.12, .42, .65, .85, 1.0))}
            sd[5] = end
            exp = next((k for k in range(1, 6) if sd[k] >= ASOF), 6)
            delayed = random.random() < p_delay
            if exp == 6:
                idx, done = (5, False) if delayed else (5, True)
            else:
                idx, done = (exp - 1, False) if delayed and exp >= 2 else (exp, False)
            comp = 5 if done else idx - 1
            hr = random.choice([0.25, 0.5, 1, 1.5, 2, 3, 4])
            ppl = random.choice([1, 1, 1, 2])
            cnt = round(U(1.2, 11) * 2000 / (hr * ppl))
            asis_mh = hr * cnt * ppl
            t_hr = round(hr * U(.2, .6), 2)
            t_ppl = ppl if random.random() < .7 else 1
            tobe_mh = t_hr * cnt * t_ppl
            found = (asis_mh - tobe_mh) / 2000
            s3, s4, s5 = found * U(.25, .45), found * U(.65, .85), found
            sec = 0
            if comp == 3: sec = s3 * U(.75, 1.05)
            elif comp == 4: sec = s4 * U(.8, 1.0)
            elif comp == 5: sec = s5 * U(.88, 1.0)
            elif idx == 3 and random.random() < .5: sec = s3 * U(.2, .5)
            sec = min(sec * mult, found)
            P = mg == "P"
            kf = None if P else abs(tobe - asis)
            ksec = None if P or sec == 0 else kf * sec / found
            krel = None if P or comp < 5 else kf * U(.8, 1.0)
            rt = rl = rx = 0
            if comp >= 4:
                rt = int(sec * U(.3, .7))
                if comp == 5: rl = int(max(0, sec - rt) * U(0, .6))
                if comp == 5 and random.random() < .3 and sec - rt - rl >= 1: rx = 1
            act_s = start + dt.timedelta(days=random.randint(-10, 30)) if start <= ASOF else None
            act_e = min(end + dt.timedelta(days=random.randint(-20, 10)), ASOF - dt.timedelta(days=3)) if done else None
            pk = random.random() < .65
            nxt = None if idx == 5 else f"S{idx+1} {STAGE[idx+1]} 착수, {sd[idx].strftime('%Y.%m')} S{idx} 완료 목표"
            if delayed and idx < 5:
                nxt = f"S{idx} 지연 만회: {random.choice(['데이터 수집 범위 재협의', '설비 인터페이스 보완', '현업 테스트 일정 재조정', '모델 정확도 개선'])} 후 {(ASOF + dt.timedelta(days=45)).strftime('%Y.%m')} 완료 예정"
            tid = f"T{len(rows) + 1:03d}"
            ramps[tid] = sd[3] - dt.timedelta(days=45)
            rows.append({
                "id": tid, "team": team, "group": group, "cat": cat, "task": task, "kind": kind, "agent": agent,
                "scope": scope, "line": line, "tgt_dept": team, "tgt_ids": ids(), "dev_dept": DEV[kind][0] if random.random() < .7 else DEV[kind][1], "dev_ids": ids(),
                "mgmt": mg, "kpi": kpi, "kpi_etc": "협상 준비시간" if kpi == "기타" else None, "kpi_etc_unit": "시간" if kpi == "기타" else None, "unit": unit,
                "stage": f"S{idx}", "next_plan": nxt, "priority": None,
                "kpi_asis": asis, "kpi_asis_basis": None if P else "2025년 1~12월 MES·현업 집계 평균", "kpi_dir": dirn,
                "kpi_tobe": tobe, "kpi_tobe_basis": None if P else "파일럿 결과 및 벤더 레퍼런스 기준",
                "kpi_found": r2(kf), "kpi_secured": r2(ksec), "kpi_secured_basis": None if ksec is None else "현장 실증 4주 측정값",
                "kpi_realized": r2(krel), "kpi_realized_basis": None if krel is None else "정식 적용 후 3개월 실적",
                "asis_hr": hr, "asis_cnt": cnt, "asis_ppl": ppl, "asis_basis": "2025년 작업일보·시스템 로그 기준",
                "tobe_hr": t_hr, "tobe_cnt": cnt, "tobe_ppl": t_ppl, "tobe_basis": "개발 사양 및 파일럿 측정 기준",
                "asis_mh": round(asis_mh, 1), "tobe_mh": round(tobe_mh, 1), "base_fte": r2(asis_mh / 2000), "found_fte": r2(found),
                "secured_fte": r2(sec) if sec else None, "secured_basis": "실증 기간 공수 측정(MES 작업시간)" if sec else None,
                "s1_date": D(sd[1]), "s2_date": D(sd[2]), "s3_date": D(sd[3]), "s3_fte": r2(s3), "s3_kpi": None if P else r2(kf * s3 / found),
                "s4_date": D(sd[4]), "s4_fte": r2(s4), "s4_kpi": None if P else r2(kf * s4 / found),
                "s5_date": D(sd[5]), "s5_fte": r2(s5), "s5_kpi": None if P else r2(kf),
                "plan_start": D(start), "plan_end": D(end), "act_start": D(act_s), "act_end": D(act_e),
                "pkms": "등록" if pk else "미등록", "pkms_name": f"{cat}-{task}" if pk else None,
                "re_team": rt or None, "re_team_date": D(sd[4] + dt.timedelta(days=30)) if rt else None,
                "re_line": rl or None, "re_line_date": D(sd[5] + dt.timedelta(days=20)) if rl else None,
                "re_etc": rx or None, "re_etc_date": D(sd[5] + dt.timedelta(days=40)) if rx else None,
            })

# 우선순위: 팀마다 발굴 FTE가 가장 큰 과제 1건은 '상', 나머지는 상 20% · 중 45% · 하 35%
for team in ORG:
    tr = sorted([r for r in rows if r["team"] == team], key=lambda r: -r["found_fte"])
    tr[0]["priority"] = "상"
    for r in tr[1:]:
        r["priority"] = random.choices(["상", "중", "하"], [.2, .45, .35])[0]

# 데이터 점검 탭을 보여 주기 위한 입력 누락 예시
by = {r["task"]: r for r in rows}
for t in ["약액 사용량 최적화", "공정 조건 이상 탐지"]:
    if by[t]["stage"] != "S5": by[t]["next_plan"] = None
next(r for r in rows if r["secured_fte"])["secured_basis"] = None
by["스마트 점검 앱"]["pkms"] = "등록"; by["스마트 점검 앱"]["pkms_name"] = None
by["단가 비교 분석 자동화"]["kpi_etc_unit"] = None

# ── 주간 스냅샷(최근 8주): 확보 FTE가 몇 차례 계단식으로 늘어 온 것으로 가정 ──────────
SLIM = ["id", "team", "group", "task", "kind", "priority", "stage", "found_fte", "secured_fte", "s3_date", "s3_fte", "s4_date", "s4_fte", "s5_date", "s5_fte", "re_team", "re_line", "re_etc", "act_end"]
steps = {}
for r in rows:
    if not r["secured_fte"]: continue
    r0 = min(ramps[r["id"]], ASOF - dt.timedelta(days=30))
    span = (ASOF - r0).days
    pts = sorted(r0 + dt.timedelta(days=random.randint(1, span)) for _ in range(random.randint(2, 4)))
    if r["team"] != "FAB운영1팀" and random.random() < .35:
        pts[-1] = ASOF   # 이번 주(기준일)에 늘어난 과제: 지난주 말 스냅샷 대비 증가로 보임
    if r["team"] == "FAB운영1팀":
        pts = [p for p in pts if p < ASOF - dt.timedelta(days=7)] or [ASOF - dt.timedelta(days=40)]
    steps[r["id"]] = pts
decrease = next(r["id"] for r in rows if r["team"] == "FAB운영1팀" and r["secured_fte"])   # 재산정으로 줄어든 과제 1건

def secured_at(r, day):
    pts = steps.get(r["id"])
    if not pts: return None
    v = r["secured_fte"] * sum(p <= day for p in pts) / len(pts)
    if r["id"] == decrease and day < ASOF: v = r["secured_fte"] * 1.25
    return r2(v) or None

history = []
for k in range(8, 0, -1):
    end = ASOF - dt.timedelta(days=ASOF.weekday() + 1) - dt.timedelta(weeks=k - 1)   # 그 주 일요일
    y, w, _ = end.isocalendar()
    history.append({"week": f"{y}-W{w:02d}", "end": D(end),
                    "tasks": [{**{k2: r[k2] for k2 in SLIM}, "secured_fte": secured_at(r, end)} for r in rows]})

# ── 출력 1: JSON ─────────────────────────────────────────────
(ROOT / "sample").mkdir(exist_ok=True)
(ROOT / "sample" / "tasks.json").write_text(json.dumps(rows, ensure_ascii=False, indent=1), encoding="utf-8")
(ROOT / "sample" / "history.json").write_text(json.dumps(history, ensure_ascii=False, indent=1), encoding="utf-8")

# ── 출력 2: 업로드 양식과 같은 컬럼 순서의 엑셀 ─────────────────
U_ = "unit"
COLS = [
 ("01 조직·과제 정보", [("* 팀","team"),("* 그룹","group"),("* 과제명(대분류)","cat"),("* 과제명(중분류)","task"),("* 과제구분","kind"),
   ("* Agent 개발 포함 여부","agent"),("* 적용범위","scope"),("* 적용범위(라인/공정)","line"),("* Target 담당 부서명","tgt_dept"),
   ("* Target 담당자 ID","tgt_ids"),("* 개발 부서명","dev_dept"),("* 개발부서 담당자 ID","dev_ids")]),
 ("02 필수 지표·KPI 선택", [("* 경영관리지표","mgmt"),("* KPI","kpi"),("* [조건] 기타 KPI","kpi_etc"),("* [조건] 기타 KPI 단위","kpi_etc_unit"),("KPI 단위 [자동]",U_)]),
 ("03 과제 진행 현황", [("* 진행단계","stage"),("* [조건] 다음 단계 추진계획","next_plan"),("우선순위","priority")]),
 ("04 최종 KPI 목표·확보·실현", [("KPI AS-IS 단위 [자동]",U_),("* [P외] KPI AS-IS 현재 수준","kpi_asis"),("* [P외] KPI AS-IS 산정근거","kpi_asis_basis"),
   ("* [P외] KPI 개선방향","kpi_dir"),("KPI TO-BE 단위 [자동]",U_),("* [P외] KPI TO-BE 목표 수준","kpi_tobe"),("* [P외] KPI TO-BE 산정근거","kpi_tobe_basis"),
   ("KPI 발굴 효과 단위 [자동]",U_),("KPI 발굴 효과 [자동]","kpi_found"),("KPI 확보 효과 단위 [자동]",U_),("KPI 확보 효과(현재까지)","kpi_secured"),
   ("* [조건] KPI 확보 효과 산정근거","kpi_secured_basis"),("KPI 실현 효과 단위 [자동]",U_),("KPI 실현 효과","kpi_realized"),("* [조건] KPI 실현 효과 산정근거","kpi_realized_basis")]),
 ("05 최종 FTE 목표·확보", [("* AS-IS 1인당 건당 소요시간(시간)","asis_hr"),("* AS-IS 연간 처리건수","asis_cnt"),("* AS-IS 투입인원","asis_ppl"),("* AS-IS 공수 산정근거","asis_basis"),
   ("* TO-BE 1인당 건당 소요시간(시간)","tobe_hr"),("* TO-BE 연간 처리건수","tobe_cnt"),("* TO-BE 투입인원","tobe_ppl"),("* TO-BE 공수 산정근거","tobe_basis"),
   ("AS-IS 연간 공수(MH) [자동]","asis_mh"),("TO-BE 연간 공수(MH) [자동]","tobe_mh"),("Baseline FTE [자동]","base_fte"),("발굴 FTE [자동]","found_fte"),
   ("확보 FTE(현재까지)","secured_fte"),("* [조건] 확보 FTE 산정근거","secured_basis")]),
 ("06 S1~S5 단계별 누적 확보 목표", [("* S1 목표일정","s1_date"),("* S2 목표일정","s2_date"),("* S3 목표일정","s3_date"),("* S3 누적 확보 목표 효과(FTE)","s3_fte"),
   ("S3 KPI 단위 [자동]",U_),("* [P외] S3 누적 확보 목표 효과(KPI)","s3_kpi"),("* S4 목표일정","s4_date"),("* S4 누적 확보 목표 효과(FTE)","s4_fte"),
   ("S4 KPI 단위 [자동]",U_),("* [P외] S4 누적 확보 목표 효과(KPI)","s4_kpi"),("* S5 목표일정","s5_date"),("* S5 누적 확보 목표 효과(FTE)","s5_fte"),
   ("S5 KPI 단위 [자동]",U_),("* [P외] S5 누적 확보 목표 효과(KPI)","s5_kpi")]),
 ("07 주요 추진 일정", [("* 계획 시작일","plan_start"),("* 계획 종료일","plan_end"),("* [조건] 실적 시작일","act_start"),("* [조건] 실적 종료일","act_end")]),
 ("08 Process-PKMS", [("* Process-PKMS 등록여부","pkms"),("* [조건] Process-PKMS 과제명","pkms_name")]),
 ("09 FTE 실현·재배치", [("팀내 업무 재배치 인원(명)","re_team"),("팀내 업무 재배치 일정","re_team_date"),("신규라인 재배치 인원(명)","re_line"),
   ("신규라인 재배치 일정","re_line_date"),("기타 실현 인원(휴직/파견 등)(명)","re_etc"),("기타 실현 일정","re_etc_date")]),
]
wb = Workbook(); ws = wb.active; ws.title = "과제"
sec_fill = PatternFill("solid", fgColor="1B365D"); hdr_fill = PatternFill("solid", fgColor="E4EAF2"); auto_fill = PatternFill("solid", fgColor="F1F3F6")
c = 1
for sec, cols in COLS:
    ws.cell(1, c, sec).font = Font(bold=True, color="FFFFFF"); ws.cell(1, c).fill = sec_fill
    if len(cols) > 1: ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + len(cols) - 1)
    for h, k in cols:
        cell = ws.cell(2, c, h); cell.font = Font(bold=True); cell.fill = auto_fill if "[자동]" in h else hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = max(12, min(30, len(h) * 1.6))
        for i, r in enumerate(rows):
            v = r.get(k)
            if isinstance(v, str) and k.endswith(("_date", "_start", "_end")) and regex.match(r"\d{4}-\d{2}-\d{2}$", v):
                ws.cell(3 + i, c, dt.date.fromisoformat(v)).number_format = "yyyy-mm-dd"
            else:
                ws.cell(3 + i, c, v)
        c += 1
ws.row_dimensions[2].height = 42
ws.freeze_panes = "E3"
wb.save(ROOT / "sample" / "AX_과제_예시.xlsx")

# ── 출력 3: index.html 예시 데이터 갱신 ──────────────────────
html = (ROOT / "index.html").read_text(encoding="utf-8")
for sid, obj in (("sample-data", rows), ("sample-history", history)):
    data = json.dumps(obj, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    html, n = regex.subn(rf'(<script id="{sid}" type="application/json">).*?(</script>)', lambda m: m.group(1) + data + m.group(2), html, flags=regex.S)
    assert n == 1, f"index.html에서 {sid} 자리를 찾지 못함"
(ROOT / "index.html").write_text(html, encoding="utf-8")

from collections import Counter
print("과제", len(rows), dict(sorted(Counter(r["stage"] for r in rows).items())), dict(Counter(r["priority"] for r in rows)))
print("발굴 FTE", round(sum(r["found_fte"] for r in rows), 1), "확보 FTE", round(sum(r["secured_fte"] or 0 for r in rows), 1))
for team in ORG:
    tr = [r for r in rows if r["team"] == team]
    now = sum(r["secured_fte"] or 0 for r in tr); prev = sum(t["secured_fte"] or 0 for t in history[-1]["tasks"] if t["team"] == team)
    print(f"  {team}: 확보 {now:.1f} (전주 {prev:.1f})")
